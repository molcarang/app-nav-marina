"""Generate the Spanish user manual as PDF and editable DOCX.

Install scripts/manual-requirements.txt into an environment, or into
artifacts/manual-tools, then run: python scripts/generate-user-manual.py
"""
from copy import deepcopy
import json
from pathlib import Path
import sys
from xml.sax.saxutils import escape

ROOT = Path(__file__).resolve().parents[1]
LOCAL_TOOLS = ROOT / "artifacts/manual-tools"
if LOCAL_TOOLS.exists():
    sys.path.insert(0, str(LOCAL_TOOLS))

import pymupdf as fitz
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.graphics import renderSVG
from reportlab.graphics.shapes import Circle, Drawing, Line, Polygon, Rect, String
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

OUT = ROOT / "docs"
ASSETS = OUT / "manual-assets"
ASSETS.mkdir(parents=True, exist_ok=True)
DATA = json.loads((OUT / "manual-content.json").read_text(encoding="utf-8"))
NAVY, TEAL, BLUE = "#143249", "#147D85", "#377DB1"
GREEN, RED, AMBER = "#21845E", "#BE4051", "#A66B18"
INK, GRAY, PALE = "#233D4B", "#59717D", "#EDF3F6"


def colour(value):
    return colors.HexColor(value)


def text(drawing, x, y, value, size=10, color=INK, bold=False, anchor="start"):
    drawing.add(String(x, y, value, fontName="Helvetica-Bold" if bold else "Helvetica",
                       fontSize=size, fillColor=colour(color), textAnchor=anchor))


def box(drawing, x, y, width, height, fill=PALE, stroke=None, radius=8):
    drawing.add(Rect(x, y, width, height, rx=radius, ry=radius,
                     fillColor=colour(fill), strokeColor=colour(stroke) if stroke else None))


def line(drawing, x1, y1, x2, y2, color=GRAY, width=1.5, dashed=False):
    drawing.add(Line(x1, y1, x2, y2, strokeColor=colour(color), strokeWidth=width,
                     strokeDashArray=[4, 3] if dashed else None))


def arrow(drawing, x1, y1, x2, y2, color=TEAL):
    import math
    line(drawing, x1, y1, x2, y2, color, 2)
    a = math.atan2(y2 - y1, x2 - x1)
    points = [x2, y2, x2 - 7 * math.cos(a - 0.45), y2 - 7 * math.sin(a - 0.45),
              x2 - 7 * math.cos(a + 0.45), y2 - 7 * math.sin(a + 0.45)]
    drawing.add(Polygon(points, fillColor=colour(color), strokeColor=None))


def ship(drawing, x, y, color=TEAL):
    drawing.add(Polygon([x, y + 10, x + 7, y - 8, x, y - 4, x - 7, y - 8],
                        fillColor=colour(color), strokeColor=colors.white, strokeWidth=1))


def badge(drawing, x, y, number):
    drawing.add(Circle(x, y, 9, fillColor=colour(TEAL), strokeColor=None))
    text(drawing, x, y - 3, str(number), 10, "#FFFFFF", True, "middle")


def make_figure(name):
    d = Drawing(500, 190)
    box(d, 0, 0, 500, 190)
    if name in ["overview", "screen"]:
        box(d, 12, 12, 476, 166, NAVY)
        text(d, 28, 157, "NavMarina", 13, "#FFFFFF", True)
        text(d, 145, 157, "CONECTADO", 9, "#76DEB6")
        text(d, 333, 157, "NOCHE   AJUSTES   [ ]", 9, "#BBD2DD")
        d.add(Circle(128, 94, 47, fillColor=colour("#1D4862"), strokeColor=colour("#8FB2C7")))
        d.add(Circle(128, 94, 34, fillColor=None, strokeColor=colour("#426B82")))
        ship(d, 128, 94, "#FFFFFF")
        text(d, 128, 127, "N", 9, "#CFE6F2", anchor="middle")
        text(d, 128, 65, "RUMBO", 8, "#CFE6F2", anchor="middle")
        labels = ["TWS", "SOG", "TWA", "COG", "SONDA", "AWA"]
        for i, label in enumerate(labels):
            x, y = 257 + (i % 3) * 72, 99 - (i // 3) * 39
            box(d, x, y, 65, 32, "#20485E", "#64889C", 4)
            text(d, x + 32, y + 12, label, 10, "#D8EAF3", True, "middle")
        box(d, 28, 25, 200, 17, "#1B5F59", radius=4)
        text(d, 37, 30, "FONDEO     /     DATOS AIS", 8, "#FFFFFF")
        text(d, 263, 46, "ESCORA  ·  TIMÓN  ·  LAT/LON", 8, "#CFE6F2")
        text(d, 263, 25, "WAYPOINT (si está activo)", 8, "#CFE6F2")
        if name == "screen":
            for x, y, n in [(229, 161, 1), (62, 100, 2), (465, 136, 3),
                            (234, 33, 4), (467, 48, 5), (467, 26, 6)]:
                badge(d, x, y, n)
    elif name == "dataflow":
        box(d, 16, 117, 126, 48, "#DCEBEF")
        text(d, 79, 146, "Sensores", 12, INK, True, "middle")
        text(d, 79, 130, "o simulador", 10, GRAY, anchor="middle")
        box(d, 186, 103, 125, 64, NAVY)
        text(d, 248, 133, "Signal K", 15, "#FFFFFF", True, "middle")
        arrow(d, 147, 140, 179, 140)
        box(d, 358, 117, 125, 48, "#D4EEE4")
        text(d, 420, 138, "NavMarina", 13, GREEN, True, "middle")
        arrow(d, 317, 140, 350, 140)
        box(d, 186, 16, 125, 44, "#E0E9F4")
        text(d, 248, 33, "OpenCPN", 13, BLUE, True, "middle")
        arrow(d, 221, 96, 221, 65, BLUE)
        arrow(d, 276, 65, 276, 96, AMBER)
        text(d, 28, 62, "Posición e instrumentos", 10, BLUE)
        text(d, 320, 38, "Destino y datos de ruta", 10, AMBER)
    elif name == "gestures":
        for x, title in [(18, "HISTORIAL"), (260, "REINICIO")]:
            box(d, x, 18, 222, 154, "#FFFFFF", "#C8D8E0")
            text(d, x + 111, 146, title, 13, NAVY, True, "middle")
            box(d, x + 65, 64, 92, 59, NAVY)
            text(d, x + 111, 94, "TWS / SOG", 11, "#FFFFFF", True, "middle")
        for x in [105, 135]:
            d.add(Circle(x, 48, 5, fillColor=colour(TEAL), strokeColor=None))
        text(d, 128, 27, "Dos toques rápidos", 11, TEAL, anchor="middle")
        text(d, 371, 39, "Mantener 3 segundos", 11, RED, True, "middle")
    elif name == "ais":
        for r in [23, 46, 69]:
            d.add(Circle(100, 91, r, fillColor=None, strokeColor=colour("#91B6C4")))
        text(d, 100, 171, "PROA", 9, TEAL, True, "middle")
        ship(d, 100, 91, NAVY)
        for x, y, n in [(115, 111, 1), (68, 123, 2), (151, 64, 3)]:
            ship(d, x, y)
            text(d, x + 10, y, str(n), 9, INK, True)
        arrow(d, 188, 90, 228, 90)
        box(d, 245, 22, 238, 146, "#FFFFFF", "#C5D8E0")
        text(d, 260, 145, "CONTACTOS · Más cercano primero", 11, NAVY, True)
        for y, label, nm in [(115, "1  Velero Alba", "0,20 NM"), (82, "2  Brisa", "0,45 NM"), (49, "3  Costa", "0,70 NM")]:
            text(d, 260, y, label, 11)
            text(d, 467, y, nm, 10, TEAL, anchor="end")
            line(d, 259, y - 11, 467, y - 11, "#D6E2E8", 0.7)
    elif name == "waypoint":
        line(d, 50, 35, 436, 151, BLUE, 3)
        text(d, 28, 17, "Inicio del tramo", 10, BLUE)
        d.add(Circle(436, 151, 6, fillColor=colour(TEAL), strokeColor=None))
        text(d, 409, 169, "Waypoint", 11, TEAL, True)
        ship(d, 244, 45, NAVY)
        text(d, 264, 34, "Barco", 10, NAVY)
        line(d, 244, 45, 231, 89, RED, 2, True)
        text(d, 182, 55, "XTE", 12, RED, True)
        arrow(d, 252, 48, 421, 143, AMBER)
        text(d, 325, 71, "Demora al destino", 10, AMBER)
        text(d, 63, 140, "XTE: separación lateral", 11, RED)
        text(d, 63, 123, "respecto al tramo", 11, RED)
    elif name == "anchor":
        cx, cy, r = 150, 94, 69
        d.add(Circle(cx, cy, r, fillColor=colour("#DAEAF1"), strokeColor=colour(TEAL), strokeDashArray=[4, 3]))
        text(d, cx, 174, "N", 12, TEAL, True, "middle")
        line(d, cx - 7, cy, cx + 7, cy, NAVY, 2)
        line(d, cx, cy - 7, cx, cy + 7, NAVY, 2)
        ship(d, 184, 117, GREEN)
        ship(d, 237, 62, RED)
        line(d, cx + 8, cy, cx + r, cy, TEAL, 1, True)
        text(d, 172, 79, "Radio", 10, TEAL)
        for y, label, color in [(141, "Cruz = centro fijado", NAVY), (111, "Verde = barco dentro", GREEN),
                                 (81, "Rojo = barco fuera", RED), (46, "El círculo no sigue al barco", GRAY)]:
            text(d, 285, y, label, 11, color)
    elif name == "anchor_states":
        states = [("Dentro", "Vigilancia activa", GREEN), ("Primera salida", "Confirmando…", AMBER),
                  ("Salida confirmada", "Aviso sonoro", RED)]
        for i, (title, label, color) in enumerate(states):
            x = 12 + i * 167
            box(d, x, 70, 143, 88, "#FFFFFF", color)
            d.add(Circle(x + 19, 135, 5, fillColor=colour(color), strokeColor=None))
            text(d, x + 12, 111, title, 11, color, True)
            text(d, x + 12, 90, label, 10)
            if i < 2:
                arrow(d, x + 146, 113, x + 164, 113, GRAY)
        text(d, 250, 43, "Confirmación: 5 s con nuevas lecturas fuera", 12, AMBER, True, "middle")
        text(d, 250, 20, "Mantener Fondeo 2 s = desactivar", 12, NAVY, True, "middle")
    elif name == "updates":
        labels = [("1", "Aviso"), ("2", "Confirmar"), ("3", "Descargar"), ("4", "Verificar"), ("5", "Abrir")]
        for i, (n, title) in enumerate(labels):
            x = 7 + i * 100
            box(d, x, 63, 86, 81, "#FFFFFF", "#A8C9D0")
            badge(d, x + 43, 121, n)
            text(d, x + 43, 87, title, 11, NAVY, True, "middle")
            if i < 4:
                arrow(d, x + 88, 102, x + 98, 102)
        text(d, 250, 30, "Se conservan los ajustes del dispositivo", 12, TEAL, True, "middle")
    return d


FIGURES = {}
for figure_name in sorted({page.get("figure") for page in DATA["pages"] if page.get("figure")}):
    figure = make_figure(figure_name)
    FIGURES[figure_name] = figure
    svg = renderSVG.drawToString(figure)
    (ASSETS / f"{figure_name}.svg").write_text(svg, encoding="utf-8")
    with fitz.open(stream=svg.encode("utf-8"), filetype="svg") as vector:
        pdf_data = vector.convert_to_pdf()
    with fitz.open(stream=pdf_data, filetype="pdf") as vector_pdf:
        vector_pdf[0].get_pixmap(matrix=fitz.Matrix(2.5, 2.5)).save(str(ASSETS / f"{figure_name}.png"))


def pdf_manual():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="ManualTitle", fontName="Helvetica-Bold", fontSize=24, leading=28,
                              textColor=colour(NAVY), spaceAfter=12))
    styles.add(ParagraphStyle(name="ManualH2", fontName="Helvetica-Bold", fontSize=11.5, leading=15,
                              textColor=colour(TEAL), spaceBefore=10, spaceAfter=4, keepWithNext=True))
    styles.add(ParagraphStyle(name="ManualBody", fontName="Helvetica", fontSize=9.6, leading=13.5,
                              textColor=colour(INK), spaceAfter=7))
    styles.add(ParagraphStyle(name="ManualIntro", parent=styles["ManualBody"], fontSize=10.4, leading=14.5,
                              textColor=colour(GRAY), spaceAfter=10))
    styles.add(ParagraphStyle(name="ManualCaption", parent=styles["ManualBody"], fontSize=8.2, leading=11,
                              textColor=colour(GRAY), spaceAfter=9))
    styles.add(ParagraphStyle(name="ManualCell", parent=styles["ManualBody"], fontSize=9, leading=12, spaceAfter=0))
    styles.add(ParagraphStyle(name="ManualCellHead", parent=styles["ManualCell"], textColor=colors.white, fontName="Helvetica-Bold"))
    styles.add(ParagraphStyle(name="ManualNote", parent=styles["ManualBody"], fontSize=9, leading=12.5,
                              backColor=colour(PALE), borderColor=colour("#C6DDE4"), borderWidth=0.6,
                              borderPadding=9, spaceBefore=9, spaceAfter=3))
    paragraphs = []
    width = A4[0] - 88

    def para(value, style="ManualBody"):
        return Paragraph(escape(value), styles[style])

    for page_number, page in enumerate(DATA["pages"], 1):
        if page_number > 1:
            paragraphs.append(PageBreak())
        paragraphs.append(para("MANUAL DE USUARIO  /  " + str(page_number).zfill(2), "ManualCaption"))
        paragraphs.append(para(page["title"], "ManualTitle"))
        if page_number == 1:
            paragraphs.append(para(DATA["edition"] + " · " + DATA["scope"], "ManualCaption"))
        paragraphs.append(para(page["intro"], "ManualIntro"))
        if page.get("figure"):
            drawing = deepcopy(FIGURES[page["figure"]])
            drawing.hAlign = "CENTER"
            paragraphs.append(drawing)
            paragraphs.append(Spacer(1, 5))
            paragraphs.append(para(page["caption"], "ManualCaption"))
        for index, step in enumerate(page.get("steps", []), 1):
            paragraphs.append(Paragraph(f"<b>{index}.</b> " + escape(step), styles["ManualBody"]))
        if page.get("table"):
            table_data = [[para(cell, "ManualCellHead") for cell in page["table"]["headers"]]]
            table_data += [[para(cell, "ManualCell") for cell in row] for row in page["table"]["rows"]]
            widths = [width * 0.77, width * 0.23] if page_number == 1 else [width * 0.32, width * 0.68]
            table = Table(table_data, colWidths=widths, repeatRows=1, hAlign="LEFT")
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colour(NAVY)),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colour("#F1F5F7")]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("LINEBELOW", (0, 0), (-1, 0), 0.5, colour(TEAL)),
            ]))
            paragraphs.extend([table, Spacer(1, 5)])
        for section in page.get("sections", []):
            paragraphs.append(para(section["heading"], "ManualH2"))
            paragraphs.append(para(section["text"]))
        if page.get("note"):
            paragraphs.append(para(page["note"], "ManualNote"))

    def frame(canvas, document):
        canvas.saveState()
        canvas.setFillColor(colour(NAVY))
        canvas.rect(0, A4[1] - 28, A4[0], 28, fill=1, stroke=0)
        canvas.setFont("Helvetica-Bold", 9)
        canvas.setFillColor(colors.white)
        canvas.drawString(44, A4[1] - 18, "NAVMARINA")
        canvas.setFont("Helvetica", 8)
        canvas.drawRightString(A4[0] - 44, A4[1] - 18, "Manual de usuario · Español")
        canvas.setStrokeColor(colour("#CADBE3"))
        canvas.line(44, 35, A4[0] - 44, 35)
        canvas.setFillColor(colour(GRAY))
        canvas.setFont("Helvetica", 8)
        canvas.drawString(44, 22, DATA["edition"])
        canvas.drawRightString(A4[0] - 44, 22, f"NavMarina  |  {document.page}")
        canvas.restoreState()

    filename = OUT / "Manual-de-usuario-NavMarina.pdf"
    document = SimpleDocTemplate(str(filename), pagesize=A4, leftMargin=44, rightMargin=44,
                                 topMargin=46, bottomMargin=48, title=DATA["title"], author="NavMarina")
    document.build(paragraphs, onFirstPage=frame, onLaterPages=frame)
    return filename


def docx_manual():
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.2677), Inches(11.6929)
    section.top_margin = section.bottom_margin = Inches(0.62)
    section.left_margin = section.right_margin = Inches(0.62)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10)
    normal.font.color.rgb = RGBColor.from_string("233D4B")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.03
    for heading, size, color in [("Heading 1", 23, "143249"), ("Heading 2", 12, "147D85")]:
        style = document.styles[heading]
        style.font.name, style.font.size = "Calibri", Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before, style.paragraph_format.space_after = Pt(8), Pt(5)
    header = section.header.paragraphs[0]
    header.text = "NAVMARINA  ·  MANUAL DE USUARIO"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string("59717D")
    footer = section.footer.paragraphs[0]
    footer.text = DATA["edition"] + "  ·  "
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.runs[0].font.size = Pt(8)
    document.core_properties.title = DATA["title"]
    document.core_properties.subject = "Uso de la consola web y Raspberry con gráficos explicativos"
    document.core_properties.author = "NavMarina"

    for number, page in enumerate(DATA["pages"], 1):
        if number > 1:
            document.add_page_break()
        p = document.add_paragraph(f"MANUAL DE USUARIO / {number:02}")
        p.runs[0].font.size = Pt(8)
        document.add_heading(page["title"], level=1)
        if number == 1:
            p = document.add_paragraph(DATA["edition"] + " · " + DATA["scope"])
            p.runs[0].font.size = Pt(8)
        document.add_paragraph(page["intro"])
        if page.get("figure"):
            document.add_picture(str(ASSETS / (page["figure"] + ".png")), width=Inches(6.6))
            p = document.add_paragraph(page["caption"])
            p.runs[0].font.size = Pt(8)
            p.runs[0].font.italic = True
        for index, step in enumerate(page.get("steps", []), 1):
            p = document.add_paragraph()
            p.add_run(f"{index}. ").bold = True
            p.add_run(step)
        if page.get("table"):
            table = document.add_table(rows=1, cols=2)
            table.autofit = False
            table.columns[0].width = Inches(2.1 if number != 1 else 5.2)
            table.columns[1].width = Inches(4.7 if number != 1 else 1.6)
            table.style = "Light Shading Accent 1"
            for cell, value in zip(table.rows[0].cells, page["table"]["headers"]):
                cell.text = value
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "143249")
                cell._tc.get_or_add_tcPr().append(shade)
                for run in cell.paragraphs[0].runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.bold = True
            for row in page["table"]["rows"]:
                for cell, value in zip(table.add_row().cells, row):
                    cell.text = value
            for row in table.rows:
                no_split = OxmlElement("w:cantSplit")
                row._tr.get_or_add_trPr().append(no_split)
                for cell in row.cells:
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after = Pt(3)
                        for run in p.runs:
                            run.font.size = Pt(9)
        for item in page.get("sections", []):
            document.add_heading(item["heading"], level=2)
            document.add_paragraph(item["text"])
        if page.get("note"):
            p = document.add_paragraph(page["note"])
            p.runs[0].font.size = Pt(9)
            p.runs[0].italic = True
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "EDF3F6")
            p._p.get_or_add_pPr().append(shade)
    filename = OUT / "Manual-de-usuario-NavMarina.docx"
    document.save(filename)
    return filename


if __name__ == "__main__":
    pdf = pdf_manual()
    docx = docx_manual()
    with fitz.open(pdf) as result:
        print(f"PDF generado: {pdf} ({len(result)} páginas)")
    print(f"Word generado: {docx}")
